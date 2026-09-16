"""Coordinate model, lossless table import, and deterministic label placement.

All stored coordinates and drafting dimensions use metres. Coordinate ordering
is chosen explicitly or read from an E/N header; coordinate magnitudes never
change the interpretation. A malformed nonempty row rejects the entire import.
"""
from __future__ import annotations

import csv
import io
import math
import os
import re
import tempfile
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence


@dataclass(frozen=True)
class Placement:
    """Manually placed CAD annotation, in unoffset metres on the WCS XY plane."""
    e: float
    n: float
    horizontal_e: float
    height: float
    arrow: float


@dataclass(frozen=True)
class Point:
    name: str
    e: float
    n: float
    z: float | None = None
    placement: Placement | None = None
    capture_id: str = ""
    group_id: str = ""

    def __iter__(self):
        return iter((self.name, self.e, self.n))


@dataclass(frozen=True)
class Settings:
    radius: float = 1.5
    leader: float = 20.0
    text_height: float = 2.0
    color: int = 1
    lineweight: int = 50
    layer: str = "COORD_POINTS"
    draw_line: bool = False
    closed: bool = False
    offset_e: float = 0.0
    offset_n: float = 0.0
    scale: float = 1.0


@dataclass
class ImportResult:
    points: list[Point]
    warnings: list[str]
    description: str = ""
    detected_order: str | None = None
    confidence: float | None = None


class CoordinateOrderRequired(ValueError):
    """Valid table data needs an explicit E/N decision before it can be used."""
    def __init__(self, reason: str, sample_rows: Sequence[Sequence[str]] = ()):
        self.reason = reason
        self.sample_rows = [list(row) for row in sample_rows[:3]]
        super().__init__(reason + "；请明确选择 EN 或 NE 后重试")


@dataclass(frozen=True)
class Label:
    name: str
    x: float
    y: float
    lx: float
    ly: float
    hx: float
    tx: float
    direction: int
    height: float | None = None
    arrow: float | None = None
    coordinate_text: tuple[str, str] | None = None


def _finite(value, label: str) -> float:
    try:
        if isinstance(value, bool):
            raise ValueError
        number = float(value)
    except (TypeError, ValueError, OverflowError) as exc:
        raise ValueError(f"{label}必须是有效数字") from exc
    if not math.isfinite(number):
        raise ValueError(f"{label}必须是有限数字，不能使用 NaN 或无穷大")
    return number


def validate_settings(settings: Settings) -> None:
    for attr, label in (("radius", "圆半径"), ("leader", "引线长度"),
                        ("text_height", "字高"), ("scale", "CAD 比例")):
        value = getattr(settings, attr)
        if not isinstance(value, (float, int)) or isinstance(value, bool):
            raise ValueError(f"{label}必须是有效数字")
        if _finite(value, label) <= 0:
            raise ValueError(f"{label}必须大于 0")
    for attr, label in (("offset_e", "东向偏移"), ("offset_n", "北向偏移")):
        value = getattr(settings, attr)
        if not isinstance(value, (float, int)) or isinstance(value, bool):
            raise ValueError(f"{label}必须是有效数字")
        _finite(value, label)
    if isinstance(settings.color, bool) or not isinstance(settings.color, int) or not 1 <= settings.color <= 255:
        raise ValueError("CAD 颜色必须是 1 至 255 的整数")
    valid_weights = {0, 5, 9, 13, 15, 18, 20, 25, 30, 35, 40, 50, 53,
                     60, 70, 80, 90, 100, 106, 120, 140, 158, 200, 211}
    if isinstance(settings.lineweight, bool) or not isinstance(settings.lineweight, int) or settings.lineweight not in valid_weights:
        raise ValueError("线宽必须是标准 CAD 线宽值（例如 25、50、70，单位为 0.01 毫米）")
    if (not isinstance(settings.layer, str) or not settings.layer.strip()
            or len(settings.layer) > 255 or settings.layer != settings.layer.strip()
            or re.search(r'[<>/\\":;?*|=\x00-\x1f\x7f]', settings.layer)):
        raise ValueError('图层名不能为空，不能超过 255 个字符，不能包含首尾空格或 <>/\\":;?*|= 等字符')
    for attr in ("draw_line", "closed"):
        if not isinstance(getattr(settings, attr), bool):
            raise ValueError(f"{attr} 必须是布尔值")
    for attr in ("radius", "leader", "text_height"):
        _finite(getattr(settings, attr) * settings.scale, "缩放后的绘图尺寸")


def _validate_point(point: Point) -> None:
    if not isinstance(point.name, str) or not point.name.strip():
        raise ValueError("点名不能为空")
    if re.search(r'[\x00-\x1f\x7f]', point.name):
        raise ValueError("点名不能包含换行、制表符等控制字符")
    _finite(point.e, f"点 {point.name} 的东 E 坐标")
    _finite(point.n, f"点 {point.name} 的北 N 坐标")
    if point.z is not None:
        _finite(point.z, f"点 {point.name} 的高程 Z")
    for attr in ("capture_id", "group_id"):
        value = getattr(point, attr)
        if (not isinstance(value, str) or len(value) > 256
                or re.search(r'[\x00-\x1f\x7f]', value) or value != value.strip()):
            raise ValueError(f"点 {point.name} 的 {attr} 元数据无效")
    if point.placement is not None:
        if not isinstance(point.placement, Placement):
            raise ValueError(f"点 {point.name} 的标注位置无效")
        for attr in ("e", "n", "horizontal_e", "height", "arrow"):
            value = getattr(point.placement, attr)
            if not isinstance(value, (int, float)) or isinstance(value, bool):
                raise ValueError(f"点 {point.name} 的标注位置必须是有效数字")
            number = _finite(value, f"点 {point.name} 的标注 {attr}")
            if attr in ("height", "arrow") and number <= 0:
                raise ValueError(f"点 {point.name} 的标注 {attr} 必须大于 0")


def transform_point(point: Point, settings: Settings) -> tuple[float, float]:
    """Apply metre offsets first, then scale to CAD drawing units."""
    validate_settings(settings)
    _validate_point(point)
    x = (point.e + settings.offset_e) * settings.scale
    y = (point.n + settings.offset_n) * settings.scale
    return _finite(x, "转换后的 CAD X"), _finite(y, "转换后的 CAD Y")


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return str(int(value))
    return str(value).strip().lstrip("\ufeff")


def _header_kind(value: str) -> str | None:
    text = re.sub(r'[（(\[]\s*(?:米|m)\s*[）)\]]|米', '', value.strip().lower())
    text = re.sub(r'[\s_\-()（）\[\]]', '', text)
    # X/Y by themselves are deliberately not mapped to east/north: surveying
    # and CAD use different conventions for these same letters.
    if text in {"e", "em", "east", "easting", "东", "东e", "e东", "东坐标", "东向坐标", "东e坐标"}:
        return "E"
    if text in {"n", "nm", "north", "northing", "北", "北n", "n北", "北坐标", "北向坐标", "北n坐标"}:
        return "N"
    if text in {"点名", "点号", "名称", "编号", "桩号", "id", "point", "pointname", "name", "点名称"}:
        return "name"
    if text in {"x", "y", "x坐标", "y坐标", "坐标x", "坐标y"}:
        return text[0] if text[0] in "xy" else text[-1]
    if text in {"z", "zm", "h", "高程", "高程z", "z高程", "标高", "高度", "elevation", "height", "z坐标", "坐标z"}:
        return "Z"
    return None


def _order_from_bounds(points: Sequence[Point], reference_bounds, settings: Settings | None,
                       sample_rows: Sequence[Sequence[str]]) -> tuple[str, float, str]:
    """Compare both interpretations against supplied CAD bounds, never magnitude.

    ``points`` contains first-coordinate/second-coordinate values in e/n. Only
    a >=80% versus <=20% separation can resolve an ambiguous table automatically.
    """
    if reference_bounds is None:
        raise CoordinateOrderRequired("没有明确 E/N 表头，也没有可用于核对的底图范围", sample_rows)
    try:
        if len(reference_bounds) != 4:
            raise ValueError
        x0, y0, x1, y1 = [_finite(value, "底图范围") for value in reference_bounds]
        if x1 <= x0 or y1 <= y0:
            raise ValueError
    except (TypeError, ValueError) as exc:
        raise ValueError("用于识别坐标顺序的底图范围无效，需提供 xmin、ymin、xmax、ymax") from exc
    setting = settings if settings is not None else Settings()
    validate_settings(setting)
    matched = {}
    for mode in ("EN", "NE"):
        count = 0
        for point in points:
            e, n = (point.e, point.n) if mode == "EN" else (point.n, point.e)
            x, y = (e+setting.offset_e)*setting.scale, (n+setting.offset_n)*setting.scale
            if math.isfinite(x) and math.isfinite(y) and x0 <= x <= x1 and y0 <= y <= y1:
                count += 1
        matched[mode] = count/len(points)
    chosen = next((mode for mode in ("EN", "NE") if matched[mode] >= .8
                   and matched["NE" if mode == "EN" else "EN"] <= .2), None)
    comparison = f"EN 匹配 {matched['EN']:.0%}，NE 匹配 {matched['NE']:.0%}"
    if chosen is None:
        raise CoordinateOrderRequired("底图范围无法唯一确定坐标顺序（"+comparison+"）", sample_rows)
    return chosen, matched[chosen], "按底图范围识别为 "+chosen+"（"+comparison+"）"


def _parse_rows(rows: Iterable[tuple[int, Sequence]], order: str, *,
                reference_bounds=None, settings: Settings | None = None) -> ImportResult:
    mode = str(order).strip().upper()
    if mode not in {"EN", "NE", "AUTO"}:
        raise ValueError("坐标顺序必须是 EN、NE 或 auto")
    points: list[Point] = []
    warnings: list[str] = []
    errors: list[str] = []
    names: dict[str, int] = {}
    # Coordinate columns are stored in their physical left-to-right order.
    # Apply the chosen EN/NE interpretation only after every data row validates.
    mapping: tuple[int | None, int, int, int | None] | None = None
    first = True
    width: int | None = None
    header_order = None
    sample_rows = []
    for line, raw in rows:
        cells = [_cell(value) for value in raw]
        if not any(cells):
            continue
        if first:
            first = False
            kinds = [_header_kind(cell) for cell in cells]
            has_en = "E" in kinds and "N" in kinds
            has_xy = "x" in kinds and "y" in kinds
            if has_en or has_xy:
                if len({kind for kind in kinds if kind}) != len([kind for kind in kinds if kind]):
                    raise ValueError(f"第 {line} 行：表头存在重复坐标列或点名列")
                name_col = kinds.index("name") if "name" in kinds else None
                coordinate_cols = sorted(i for i, kind in enumerate(kinds) if kind in {"E", "N", "x", "y"})
                if len(coordinate_cols) != 2:
                    raise ValueError(f"第 {line} 行：需要恰好两列平面坐标")
                a_col, b_col = coordinate_cols
                z_col = kinds.index("Z") if "Z" in kinds else None
                if has_en:
                    header_order = "EN" if kinds.index("E") < kinds.index("N") else "NE"
                    if mode != "AUTO" and mode != header_order:
                        warnings.append(f"第 {line} 行：表头与所选 {mode} 顺序不一致，已按手动选择的 {mode} 导入")
                mapping = name_col, a_col, b_col, z_col
                width = len(cells)
                used = {a_col, b_col} | ({name_col} if name_col is not None else set()) | ({z_col} if z_col is not None else set())
                extra = [cells[i] or f"第 {i+1} 列" for i in range(len(cells)) if i not in used]
                if extra:
                    warnings.append("仅导入点名、平面坐标和高程，未导入字段：" + "、".join(extra))
                continue
        try:
            if mapping is None:
                if len(cells) not in {2, 3, 4}:
                    raise ValueError(f"需要 2 列坐标、3 列（点名、坐标、坐标）或 4 列（另含高程），实际 {len(cells)} 列；空字段不能省略")
                if width is not None and len(cells) != width:
                    raise ValueError(f"列数与前面的 {width} 列不一致")
                width = len(cells)
                name_col = 0 if len(cells) in {3, 4} else None
                first_coord = 1 if name_col is not None else 0
                a_col, b_col = first_coord, first_coord + 1
                z_col = 3 if len(cells) == 4 else None
            else:
                name_col, a_col, b_col, z_col = mapping
                if len(cells) != width:
                    raise ValueError(f"列数与表头的 {width} 列不一致，实际 {len(cells)} 列")
            name = cells[name_col] if name_col is not None else f"点{len(points)+1}"
            if not cells[a_col] or not cells[b_col]:
                raise ValueError("东 E 或北 N 坐标为空")
            z = _finite(cells[z_col], "高程 Z") if z_col is not None and cells[z_col] else None
            point = Point(name, _finite(cells[a_col], "第一列平面坐标"), _finite(cells[b_col], "第二列平面坐标"), z)
            _validate_point(point)
            if name in names:
                warnings.append(f"第 {line} 行：点名“{name}”与第 {names[name]} 行重复，已保留两个点")
            else:
                names[name] = line
            points.append(point)
            if len(sample_rows) < 3:
                sample_rows.append(cells)
        except ValueError as exc:
            errors.append(f"第 {line} 行：{exc}")
    if errors:
        shown = errors[:20]
        if len(errors) > 20:
            shown.append(f"另有 {len(errors)-20} 行错误")
        raise ValueError("导入已取消，原有坐标未改变：\n" + "\n".join(shown))
    if not points:
        raise ValueError("没有找到坐标数据，请提供点名、东 E、北 N，或两列坐标")
    confidence = None
    if mode != "AUTO":
        chosen, description = mode, "按手动选择的 "+mode+" 顺序导入"
    elif header_order is not None:
        chosen, confidence, description = header_order, 1.0, "按明确 E/N 表头识别为 "+header_order
    else:
        chosen, confidence, description = _order_from_bounds(points, reference_bounds, settings, sample_rows)
    if chosen == "NE":
        points = [Point(point.name, point.n, point.e, point.z) for point in points]
    if any(point.z is not None for point in points):
        description += "；已保留高程 Z（二维落图使用 E/N）"
    if confidence is not None and confidence < 1:
        warnings.append(f"所选 {chosen} 顺序仍有 {1-confidence:.0%} 的点在底图范围外，请核对点表和坐标系")
    return ImportResult(points, warnings, description, chosen, confidence)


def parse_text(text: str, order: str = "EN", *, reference_bounds=None,
               settings: Settings | None = None) -> ImportResult:
    """Read comma/tab/semicolon CSV or whitespace columns, preserving blanks.

    Unheaded three-column data is name/coordinate/coordinate, including numeric
    names; a fourth column is optional elevation Z. Two-coordinate data receives
    sequential names. Auto follows E/N headers or a decisive drawing-bounds
    comparison; otherwise CoordinateOrderRequired requests a human decision.
    """
    if not isinstance(text, str):
        raise TypeError("输入必须是文本")
    text = text.lstrip("\ufeff")
    lines = text.splitlines(keepends=True)
    first_number = next((number for number, line in enumerate(lines) if line.strip()), None)
    first_line = lines[first_number].rstrip("\r\n") if first_number is not None else ""
    sep_match = re.fullmatch(r"\s*sep=([,;\t，])\s*", first_line, re.IGNORECASE)
    explicit_delimiter = sep_match.group(1) if sep_match else None
    if sep_match:
        # Preserve original line numbering while discarding Excel's delimiter
        # directive, which is metadata rather than a point row.
        lines[first_number] = "\n"
        text = "".join(lines)
    candidates = [candidate for candidate in ("\t", ",", ";", "，") if candidate in first_line]
    delimiter = explicit_delimiter
    if candidates and delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(first_line, delimiters="".join(candidates)).delimiter
        except csv.Error:
            # Still let the strict full reader report malformed quoting.
            delimiter = candidates[0]
    if delimiter:
        reader = csv.reader(io.StringIO(text), delimiter=delimiter, strict=True)
        rows = []
        previous_end = 0
        try:
            for row in reader:
                rows.append((previous_end + 1, row))
                previous_end = reader.line_num
        except csv.Error as exc:
            raise ValueError(f"第 {reader.line_num} 行：CSV 格式错误：{exc}") from exc
    else:
        rows = [(number, line.split()) for number, line in enumerate(text.splitlines(), 1)]
    return _parse_rows(rows, order, reference_bounds=reference_bounds, settings=settings)


def read_points(path: str | Path, order: str = "EN", *, reference_bounds=None,
                settings: Settings | None = None) -> ImportResult:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in {".txt", ".csv", ".tsv"}:
        data = path.read_bytes()
        encodings = ("utf-16",) if data.startswith((b"\xff\xfe", b"\xfe\xff")) else ("utf-8-sig", "gb18030")
        for encoding in encodings:
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("文本编码无法识别，请另存为 UTF-8、GB18030 或带 BOM 的 UTF-16 文件")
        return parse_text(text, order, reference_bounds=reference_bounds, settings=settings)
    if suffix == ".xlsx":
        try:
            import openpyxl
        except ImportError as exc:
            raise ValueError("读取 .xlsx 需要 openpyxl，请运行依赖安装程序") from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = workbook.worksheets[0]
            rows = list(sheet.iter_rows(values_only=True))
            # Ignore empty formatting-only columns, but never compress blanks
            # between actual columns. Formula cells fail numeric validation.
            last_column = max((max((i+1 for i, value in enumerate(row) if value is not None), default=0) for row in rows), default=0)
            result = _parse_rows(((i, row[:last_column]) for i, row in enumerate(rows, 1)), order,
                                 reference_bounds=reference_bounds, settings=settings)
            if len(workbook.worksheets) > 1:
                result.warnings.append(f"工作簿有 {len(workbook.worksheets)} 个工作表，仅导入第一个：{sheet.title}")
            return result
        finally:
            workbook.close()
    if suffix == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise ValueError("读取旧版 .xls 需要 xlrd；也可以在 Excel 中另存为 .xlsx") from exc
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        try:
            sheet = workbook.sheet_by_index(0)
            result = _parse_rows(((i+1, sheet.row_values(i)) for i in range(sheet.nrows)), order,
                                 reference_bounds=reference_bounds, settings=settings)
            if workbook.nsheets > 1:
                result.warnings.append(f"工作簿有 {workbook.nsheets} 个工作表，仅导入第一个：{sheet.name}")
            return result
        finally:
            workbook.release_resources()
    raise ValueError("不支持此坐标文件格式，请选择 TXT、CSV、TSV、XLSX 或 XLS")


def write_points(path: str | Path, points: Sequence[Point]) -> None:
    """Atomically save source E/N metres without offsets or rounding."""
    path = Path(path)
    for point in points:
        _validate_point(point)
    if path.suffix.lower() not in {".csv", ".txt", ".tsv", ".xlsx"}:
        raise ValueError("坐标导出支持 CSV、TXT、TSV 或 XLSX")
    with tempfile.NamedTemporaryFile(prefix=".coord-", suffix=path.suffix, dir=path.parent, delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        _write_points_file(temporary_path, points)
        os.replace(temporary_path, path)
    finally:
        temporary_path.unlink(missing_ok=True)


def _write_points_file(path: Path, points: Sequence[Point]) -> None:
    has_z = any(point.z is not None for point in points)
    header = ("点名", "东E", "北N", "高程Z") if has_z else ("点名", "东E", "北N")
    def row(point):
        return (point.name, point.e, point.n, point.z) if has_z else (point.name, point.e, point.n)
    if path.suffix.lower() == ".xlsx":
        try:
            import openpyxl
        except ImportError as exc:
            raise ValueError("导出 .xlsx 需要 openpyxl") from exc
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "坐标"
        sheet.append(header)
        for point in points:
            sheet.append(row(point))
            sheet.cell(sheet.max_row, 1).data_type = "s"
        sheet.column_dimensions["A"].width = 24
        sheet.column_dimensions["B"].width = 22
        sheet.column_dimensions["C"].width = 22
        if has_z:
            sheet.column_dimensions["D"].width = 22
        sheet.freeze_panes = "A2"
        workbook.save(path)
        workbook.close()
        return
    if path.suffix.lower() not in {".csv", ".txt", ".tsv"}:
        raise ValueError("坐标导出支持 CSV、TXT、TSV 或 XLSX")
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, delimiter="\t" if path.suffix.lower() in {".txt", ".tsv"} else ",")
        writer.writerow(header)
        writer.writerows(row(point) for point in points)


def _tangent(i: int, points: Sequence[Point], closed: bool) -> float:
    if len(points) < 2:
        return 0.0
    p = points[i]
    before = points[i-1] if i or closed else p
    after = points[(i+1) % len(points)] if i+1 < len(points) or closed else p
    vectors = [(p.e-before.e, p.n-before.n), (after.e-p.e, after.n-p.n)]
    units = [(dx / length, dy / length) for dx, dy in vectors if (length := math.hypot(dx, dy)) > 0]
    if not units:
        return 0.0
    dx, dy = sum(v[0] for v in units), sum(v[1] for v in units)
    if math.hypot(dx, dy) < 1e-10:
        dx, dy = units[-1]
    return math.atan2(dy, dx)


def _intersects(a, b) -> bool:
    ax, ay, bx, by = a
    cx, cy, dx, dy = b
    ux, uy, vx, vy = bx-ax, by-ay, dx-cx, dy-cy
    denominator = ux*vy-uy*vx
    if abs(denominator) < 1e-12:
        return False
    t = ((cx-ax)*vy-(cy-ay)*vx)/denominator
    u = ((cx-ax)*uy-(cy-ay)*ux)/denominator
    return 0.02 < t < 0.98 and 0.02 < u < 0.98


def _bounds(segment):
    ax, ay, bx, by = segment
    return min(ax, bx), min(ay, by), max(ax, bx), max(ay, by)


class _SpatialIndex:
    """Uniform grid keeps ordinary label checks local; worst case is quadratic."""
    def __init__(self, size: float):
        self.size = size
        self.cells = defaultdict(list)
        self.overflow: list[int] = []
        self.values: list[tuple] = []

    def _keys(self, box):
        x1, y1, x2, y2 = (math.floor(value/self.size) for value in box)
        if (x2-x1+1)*(y2-y1+1) > 256:
            return None
        return ((x, y) for x in range(x1, x2+1) for y in range(y1, y2+1))

    def add(self, box, value):
        index = len(self.values)
        self.values.append(value)
        keys = self._keys(box)
        if keys is None:
            self.overflow.append(index)
        else:
            for key in keys:
                self.cells[key].append(index)

    def query(self, box):
        keys = self._keys(box)
        if keys is None:
            return self.values
        found = set(self.overflow)
        for key in keys:
            found.update(self.cells.get(key, ()))
        return [self.values[index] for index in sorted(found)]


def calc_labels(points: Sequence[Point], settings: Settings) -> list[Label]:
    """Place alternating labels with local overlap/crossing penalties.

    Placement is performed in metres and transformed only at the end, so a CAD
    scale or translation never changes leader choices. A spatial grid replaces
    global repeated collision passes. Dense/coincident points can still overlap;
    this is a best-effort layout, not a geometric non-overlap guarantee.
    """
    validate_settings(settings)
    if not points:
        return []
    for point in points:
        _validate_point(point)
    th, length = settings.text_height, settings.leader
    rects = _SpatialIndex(max(length*1.5, th*16))
    segments = _SpatialIndex(max(length*1.5, th*16))
    result = []
    directions = (0, 90, 180, 270, 45, 135, 225, 315, 30, 60, 120, 150, 210, 240, 300, 330)
    scale = settings.scale
    x_transform = lambda value: _finite((value+settings.offset_e)*scale, "转换后的 CAD X")
    y_transform = lambda value: _finite((value+settings.offset_n)*scale, "转换后的 CAD Y")
    for i, point in enumerate(points):
        if point.placement is not None:
            placement = point.placement
            direction = 1 if placement.horizontal_e >= placement.e else -1
            height = _finite(placement.height*scale, "转换后的标注字高")
            arrow = _finite(placement.arrow*scale, "转换后的箭头长度")
            result.append(Label(point.name, x_transform(point.e), y_transform(point.n),
                                x_transform(placement.e), y_transform(placement.n),
                                x_transform(placement.horizontal_e), x_transform(placement.e),
                                direction, height, arrow,
                                (f"E={point.e:.3f}", f"N={point.n:.3f}")))
            box = (min(placement.e, placement.horizontal_e), placement.n-placement.height*2.7,
                   max(placement.e, placement.horizontal_e), placement.n+placement.height*1.3)
            segment = point.e, point.n, placement.e, placement.n
            rects.add(box, box)
            segments.add(_bounds(segment), segment)
            continue
        # Nearby entries preserve the original tool's polyline locality without
        # an all-pairs nearest-neighbour calculation.
        near = min((math.hypot(point.e-points[j].e, point.n-points[j].n)
                    for j in range(max(0, i-20), min(len(points), i+21)) if j != i), default=math.inf)
        base = _tangent(i, points, settings.closed) + math.pi/2
        if math.cos(base) < -0.01 or (abs(math.cos(base)) <= 0.01 and math.sin(base) < 0):
            base += math.pi
        if i % 2:
            base += math.pi
        base_deg = math.degrees(base) % 360
        ordered = sorted(directions, key=lambda angle: abs((angle-base_deg+180) % 360-180))
        effective = max(settings.radius+th, length * (1.4 if near < th*2.5 else 1.25 if near < th*5 else 1.15 if near < th*8 else 1))
        lengths = (effective, max(effective, min(effective*1.15, length*1.6))) if near < th*4 else (effective,)
        width = max(th, sum(2 if ord(char) > 127 else 1 for char in point.name)*th*0.8)
        best = None
        for candidate_length in lengths:
            for rank, angle in enumerate(ordered):
                rad = math.radians(angle)
                direction = 1 if math.cos(rad) >= -0.01 else -1
                lx, ly = point.e+candidate_length*math.cos(rad), point.n+candidate_length*math.sin(rad)
                hx = lx+direction*width
                margin = th*0.3
                box = min(lx, hx)-margin, ly-margin, max(lx, hx)+margin, ly+th*1.5
                segment = point.e, point.n, lx, ly
                sb = _bounds(segment)
                score = rank*0.3 + min(angle % 90, 90-angle % 90)*0.5 + (candidate_length-effective)*4
                for other in rects.query((min(box[0], sb[0]), min(box[1], sb[1]), max(box[2], sb[2]), max(box[3], sb[3]))):
                    ox1, oy1, ox2, oy2 = other
                    score += max(0, min(box[2], ox2)-max(box[0], ox1))*max(0, min(box[3], oy2)-max(box[1], oy1))*10
                    for t in (0.25, 0.5, 0.75):
                        mx, my = point.e+(lx-point.e)*t, point.n+(ly-point.n)*t
                        if ox1 <= mx <= ox2 and oy1 <= my <= oy2:
                            score += 300
                score += sum(200 for other in segments.query(sb) if _intersects(segment, other))
                if best is None or score < best[0]:
                    best = score, lx, ly, hx, direction, box, segment
                if score <= 0:
                    break
            if best[0] <= 0:
                break
        _, lx, ly, hx, direction, box, segment = best
        rects.add(box, box)
        segments.add(_bounds(segment), segment)
        result.append(Label(point.name, x_transform(point.e), y_transform(point.n),
                            x_transform(lx), y_transform(ly), x_transform(hx),
                            x_transform(lx if direction == 1 else hx), direction))
    return result


def annotation_segments(label: Label, settings: Settings) -> list[tuple[tuple[float, float], tuple[float, float]]]:
    """Return CAD XY line pairs; ordinary annotation circles are separate.

    Captured labels have a hollow V with each wing of length ``label.arrow``.
    Its wings point towards the elbow, 30 degrees either side of that vector.
    """
    x, y = label.x, label.y
    if label.arrow is None:
        radius = settings.radius*settings.scale
        result = [((x-radius, y), (x+radius, y)), ((x, y-radius), (x, y+radius))]
    else:
        angle = math.atan2(label.ly-y, label.lx-x)
        result = [((x, y), (x+label.arrow*math.cos(angle+delta),
                            y+label.arrow*math.sin(angle+delta)))
                  for delta in (-math.pi/6, math.pi/6)]
    result.extend([((x, y), (label.lx, label.ly)),
                   ((label.lx, label.ly), (label.hx, label.ly))])
    for segment in result:
        for vertex in segment:
            for value in vertex:
                _finite(value, "标注线段坐标")
    return result


def annotation_texts(label: Label, settings: Settings) -> list[tuple[str, float, float, float, str]]:
    """Return (text, CAD x, CAD baseline y, CAD height, left/right alignment)."""
    height = label.height if label.height is not None else settings.text_height*settings.scale
    align = "right" if label.direction < 0 else "left"
    if label.coordinate_text is None:
        result = [(label.name, label.lx, label.ly+height*.05, height, align)]
    else:
        texts = (label.name, *label.coordinate_text)
        result = [(text, label.lx, label.ly+offset*height, height, align)
                  for text, offset in zip(texts, (.3, -1.2, -2.7))]
    for _, x, y, height, _ in result:
        for value in (x, y, height):
            _finite(value, "标注文字坐标或字高")
    return result


def connection_groups(points: Sequence[Point], labels: Sequence[Label]) -> list[list[Label]]:
    """Return separate road sequences in source order, never joining road IDs.

    Ordinary ungrouped points keep their historical order within contiguous
    runs. A captured road may be resumed later; its stable group ID joins only
    that road's own points. Callers apply draw_line/closed to each group.
    """
    if len(points) != len(labels):
        raise ValueError("坐标和标注数量不一致")
    groups: list[list[Label]] = []
    roads: dict[str, list[Label]] = {}
    ordinary = None
    for point, label in zip(points, labels):
        if point.group_id:
            ordinary = None
            if point.group_id not in roads:
                roads[point.group_id] = []
                groups.append(roads[point.group_id])
            roads[point.group_id].append(label)
        else:
            if ordinary is None:
                ordinary = []
                groups.append(ordinary)
            ordinary.append(label)
    return groups
